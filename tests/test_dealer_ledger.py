import sqlite3
from datetime import date

import pytest
from openpyxl import load_workbook

from models.database import Database
from reports.dealer_ledger_export import export_dealer_ledger_excel, export_dealer_ledger_pdf
from services.dealer_ledger_service import DealerLedgerService, LedgerError


@pytest.fixture
def ledger(tmp_path):
    database = Database(tmp_path / "ledger.sqlite3")
    database.initialize()
    service = DealerLedgerService(database)
    dealer_id = service.create_dealer("Birat Pharma", company="Birat Distributors", phone="9800000000")
    return database, service, dealer_id


def test_ledger_running_balance_and_date_range_opening(ledger):
    _database, service, dealer_id = ledger
    service.add_invoice(dealer_id, "2026-10-01", "INV-001", 125000, due_date="2026-10-15", remark="Stock")
    service.add_payment(dealer_id, "2026-10-05", 25000, payment_method="bank", remark="Part payment")
    service.add_invoice(dealer_id, "2026-10-10", "INV-002", 50000, remark="Urgent order")

    report = service.ledger(dealer_id, "2026-10-05", "2026-10-31")

    assert report["opening_balance_paisa"] == 125000
    assert report["total_debit_paisa"] == 50000
    assert report["total_credit_paisa"] == 25000
    assert report["closing_balance_paisa"] == 150000
    assert [row["balance_paisa"] for row in report["entries"]] == [100000, 150000]
    assert report["entries"][0]["display_payment_date"] == "2026-10-05"


def test_duplicate_invoice_number_is_rejected_per_dealer(ledger):
    _database, service, dealer_id = ledger
    service.add_invoice(dealer_id, "2026-10-01", "INV-001", 10000)
    with pytest.raises(LedgerError, match="already exists"):
        service.add_invoice(dealer_id, "2026-10-02", "inv-001", 20000)


def test_invalid_date_range_and_amounts_are_rejected(ledger):
    _database, service, dealer_id = ledger
    with pytest.raises(LedgerError, match="greater than zero"):
        service.add_invoice(dealer_id, "2026-10-01", "INV-0", 0)
    with pytest.raises(LedgerError, match="on or before"):
        service.ledger(dealer_id, "2026-10-31", "2026-10-01")


def test_due_date_and_payment_method_validation(ledger):
    _database, service, dealer_id = ledger
    with pytest.raises(LedgerError, match="before the invoice"):
        service.add_invoice(dealer_id, "2026-10-10", "INV-DUE", 10000, due_date="2026-10-09")
    with pytest.raises(LedgerError, match="payment method"):
        service.add_payment(dealer_id, "2026-10-10", 10000, payment_method="unknown")


def test_linked_payment_cannot_cross_dealers_or_overpay(ledger):
    database, service, dealer_id = ledger
    service.add_invoice(dealer_id, "2026-10-01", "OWN-INV", 10000)
    other_dealer = service.create_dealer("Other wholesaler")
    with pytest.raises(LedgerError, match="not found"):
        service.add_payment(dealer_id, "2026-10-02", 1000, linked_invoice_no="OTHER-INV")
    service.add_invoice(other_dealer, "2026-10-01", "OTHER-INV", 10000)
    service.add_payment(dealer_id, "2026-10-02", 7000, linked_invoice_no="OWN-INV")
    with pytest.raises(LedgerError, match="exceeds"):
        service.add_payment(dealer_id, "2026-10-03", 4000, linked_invoice_no="OWN-INV")
    assert service.ledger(dealer_id)["closing_balance_paisa"] == 3000


def test_ledger_pagination_preserves_running_balance(ledger):
    database, service, dealer_id = ledger
    with database.connect() as connection:
        connection.executemany(
                """INSERT INTO ledger_entries (party_type, party_id, dealer_id, entry_date, invoice_no, debit_paisa)
                    VALUES ('dealer', ?, ?, '2026-10-01', ?, 100)""",
                [(dealer_id, dealer_id, f"PAGE-{number:03d}") for number in range(300)],
        )

    first_page = service.ledger(dealer_id, limit=250, offset=0)
    second_page = service.ledger(dealer_id, limit=250, offset=250)

    assert len(first_page["entries"]) == 250
    assert first_page["entries"][-1]["balance_paisa"] == 25000
    assert len(second_page["entries"]) == 50
    assert second_page["page_opening_balance_paisa"] == 25000
    assert second_page["entries"][0]["balance_paisa"] == 25100
    assert second_page["closing_balance_paisa"] == 30000


def test_v1_database_migrates_without_losing_dealer(tmp_path):
    path = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(path) as connection:
        connection.execute(
            "CREATE TABLE dealers (id INTEGER PRIMARY KEY, name TEXT NOT NULL, company TEXT NOT NULL DEFAULT '', phone TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)"
        )
        connection.execute("INSERT INTO dealers (name, company, phone) VALUES ('Legacy Dealer', 'Old Co', '9811111111')")
        connection.execute("PRAGMA user_version = 1")

    database = Database(path)
    database.initialize()

    with database.connect() as connection:
        dealer = connection.execute("SELECT * FROM dealers WHERE id = 1").fetchone()
        version = connection.execute("PRAGMA user_version").fetchone()[0]
        ledger_table = connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='ledger_entries'").fetchone()
        language = connection.execute("SELECT language FROM shop_settings WHERE id = 1").fetchone()[0]
    assert dealer["name"] == "Legacy Dealer"
    assert dealer["company"] == "Old Co"
    assert dealer["contact_person"] == ""
    assert version == 7
    assert ledger_table is not None
    assert language == "en"


def test_search_indexes_exist_after_migrations(tmp_path):
    database = Database(tmp_path / "indexed.sqlite3")
    database.initialize()
    with database.connect() as connection:
        version = connection.execute("PRAGMA user_version").fetchone()[0]
        indexes = {row["name"] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='index'")}
    assert version == 7
    assert {"idx_purchases_invoice_number", "idx_dealers_name", "idx_dealers_company"} <= indexes


def test_excel_and_pdf_exports_have_paper_ledger_columns(ledger, tmp_path):
    _database, service, dealer_id = ledger
    service.update_dealer(dealer_id, name="<script>alert(1)</script>")
    service.add_invoice(dealer_id, "2026-10-01", '=HYPERLINK("https://example.invalid")', 125000,
                        remark="<script>alert(1)</script>")
    service.add_payment(dealer_id, "2026-10-03", 25000, payment_method="cash", remark="Cash paid")
    report = service.ledger(dealer_id, "2026-10-01", "2026-10-31")

    excel_path = export_dealer_ledger_excel(report, tmp_path / "dealer-ledger.xlsx")
    pdf_path = export_dealer_ledger_pdf(report, tmp_path / "dealer-ledger.pdf")

    sheet = load_workbook(excel_path, data_only=True)["Dealer Ledger"]
    assert [sheet.cell(7, column).value for column in range(1, 8)] == [
        "Date", "Invoice No.", "Debit", "Date", "Credit", "Remarks", "Balance"
    ]
    assert sheet.cell(8, 2).value.startswith("'=")
    assert sheet.cell(8, 2).data_type == "s"
    assert sheet.cell(8, 6).value == "<script>alert(1)</script>"
    assert sheet.cell(9, 5).value == 250
    assert pdf_path.read_bytes().startswith(b"%PDF")
