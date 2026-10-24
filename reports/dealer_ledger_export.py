from __future__ import annotations

from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HEADERS = ["Date", "Invoice No.", "Debit", "Date", "Credit", "Remarks", "Balance"]
TEAL = "087F78"


def _amount(paisa: int) -> float:
    return paisa / 100


def _excel_text(value: Any) -> str:
    text = str(value or "")
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _excel_row(entry: dict[str, Any]) -> list[Any]:
    return [
        _excel_text(entry["display_date"]),
        _excel_text(entry["invoice_no"]),
        _amount(entry["debit_paisa"]),
        _excel_text(entry["display_payment_date"]),
        _amount(entry["credit_paisa"]),
        _excel_text(entry["remark"]),
        _amount(entry["balance_paisa"]),
    ]


def export_dealer_ledger_excel(report: dict[str, Any], path: str | Path) -> Path:
    target = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dealer Ledger"
    sheet.merge_cells("A1:G1")
    sheet["A1"] = "DEALER LEDGER"
    sheet["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=TEAL)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:G2")
    dealer = report["dealer"]
    sheet["A2"] = _excel_text(f"{dealer['name']}  |  {dealer.get('company', '')}  |  {dealer.get('phone', '')}")
    sheet["A2"].font = Font(bold=True, size=11)
    sheet.merge_cells("A3:G3")
    sheet["A3"] = f"Period: {report['start_date']} to {report['end_date']}"
    sheet["A5"] = f"Opening balance: Rs. {_amount(report['opening_balance_paisa']):,.2f}"
    header_row = 7
    for column, label in enumerate(HEADERS, 1):
        cell = sheet.cell(header_row, column, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    thin_teal = Side(style="thin", color="A9C8C4")
    for row_index, entry in enumerate(report["entries"], header_row + 1):
        for column, value in enumerate(_excel_row(entry), 1):
            cell = sheet.cell(row_index, column, value)
            cell.border = Border(bottom=thin_teal)
            if column in (3, 5, 7):
                cell.number_format = '"Rs. "#,##0.00;[Red]-"Rs. "#,##0.00'
                cell.alignment = Alignment(horizontal="right")
    totals_row = header_row + len(report["entries"]) + 1
    sheet.cell(totals_row, 2, "TOTAL")
    sheet.cell(totals_row, 3, _amount(report["total_debit_paisa"]))
    sheet.cell(totals_row, 5, _amount(report["total_credit_paisa"]))
    sheet.cell(totals_row, 7, _amount(report["closing_balance_paisa"]))
    for column in (2, 3, 5, 7):
        cell = sheet.cell(totals_row, column)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E1EFED")
        if column in (3, 5, 7):
            cell.number_format = '"Rs. "#,##0.00;[Red]-"Rs. "#,##0.00'
    widths = [14, 18, 17, 14, 17, 46, 19]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A8"
    sheet.auto_filter.ref = f"A{header_row}:G{max(header_row, totals_row - 1)}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(target)
    return target


def export_dealer_ledger_pdf(report: dict[str, Any], path: str | Path) -> Path:
    target = Path(path)
    dealer = report["dealer"]
    dealer_name = escape(str(dealer["name"]))
    dealer_company = escape(str(dealer.get("company", "")))
    dealer_phone = escape(str(dealer.get("phone", "")))
    document = SimpleDocTemplate(
        str(target), pagesize=landscape(A4),
        rightMargin=22, leftMargin=22, topMargin=20, bottomMargin=22,
        title=f"Dealer Ledger - {dealer_name}",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("LedgerTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=15,
                                 leading=18, alignment=TA_CENTER, textColor=colors.HexColor(f"#{TEAL}"), spaceAfter=5)
    detail_style = ParagraphStyle("LedgerDetail", parent=styles["Normal"], fontSize=8, leading=10, alignment=TA_LEFT)
    story = [
        Paragraph("DIPESH MEDICAL HALL  |  DEALER LEDGER", title_style),
        Paragraph(f"Dealer: {dealer_name} &nbsp;&nbsp; {dealer_company} &nbsp;&nbsp; {dealer_phone}", detail_style),
        Paragraph(f"Period: {report['start_date']} to {report['end_date']} &nbsp;&nbsp;&nbsp; Opening balance: Rs. {_amount(report['opening_balance_paisa']):,.2f}", detail_style),
        Spacer(1, 9),
    ]
    data: list[list[Any]] = [HEADERS]
    for entry in report["entries"]:
        data.append([
            entry["display_date"], entry["invoice_no"] or "",
            f"{_amount(entry['debit_paisa']):,.2f}" if entry["debit_paisa"] else "",
            entry["display_payment_date"],
            f"{_amount(entry['credit_paisa']):,.2f}" if entry["credit_paisa"] else "",
            Paragraph(escape(str(entry["remark"] or "")), detail_style),
            f"{_amount(entry['balance_paisa']):,.2f}",
        ])
    data.append([
        "", "TOTAL", f"{_amount(report['total_debit_paisa']):,.2f}", "",
        f"{_amount(report['total_credit_paisa']):,.2f}", "CLOSING BALANCE",
        f"{_amount(report['closing_balance_paisa']):,.2f}",
    ])
    table = Table(data, colWidths=[52, 72, 67, 52, 67, 390, 74], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{TEAL}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#A9C8C4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (2, 1), (2, -1), "RIGHT"),
        ("ALIGN", (4, 1), (4, -1), "RIGHT"),
        ("ALIGN", (6, 1), (6, -1), "RIGHT"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E1EFED")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    document.build(story)
    return target
